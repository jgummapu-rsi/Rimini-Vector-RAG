"""Ingest-side API routes: upload, job/document lifecycle, and ops endpoints."""

from __future__ import annotations

import csv
import hashlib
import io
import logging
import mimetypes
import os
import re
from urllib.parse import quote

import pypdfium2 as pdfium
import xlrd
from docx import Document as Docx
from docx.table import Table
from docx.text.paragraph import Paragraph
from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile, status
from fastapi.responses import Response
from openpyxl import load_workbook
from PIL import Image, ImageOps
from starlette.concurrency import run_in_threadpool
from striprtf.striprtf import rtf_to_text

from app.api._common import (
    _owned_or_404,
    _read_capped,
    _resolve_target_document,
    _visible_or_404,
)
from app.api.admission import admit_request
from app.api.auth import (
    get_container,
    get_principal,
    require_delete,
    require_ingest,
    require_operator,
)
from app.ingest.pipeline.provenance import location_str
from app.ingest.pipeline.runner import invalidate_cache_for
from app.retrieval.rag.access import can_view
from app.retrieval.rag.provenance import citation_provenance
from app.shared.container import Container
from app.shared.domain.models import (
    EXT_TO_SOURCE_TYPE,
    Document,
    DocumentVersion,
    Job,
    JobStage,
    JobStatus,
    Principal,
    Role,
    Scope,
    Visibility,
)
from app.shared.ids import new_object_id
from app.shared.observability import bind

router = APIRouter()
log = logging.getLogger("api")


@router.get("/healthz")
def healthz(container: Container = Depends(get_container)) -> dict:
    """Liveness + active-backend summary, no auth required."""
    return {
        "status": "ok",
        "backends": {
            "metadata": "postgres",
            "blob": "localfs",
            "vector": "pgvector",
            "queue": "postgres",
        },
    }


@router.get("/metrics")
def metrics(
    principal: Principal = Depends(require_operator),
    container: Container = Depends(get_container),
) -> dict:
    """Counters + timings, shared across API and worker processes.

    Admin-only: these counters are deployment-wide, not tenant-scoped, so they
    describe every tenant's ingest/query volume at once. They were previously
    readable with no credentials at all.
    """
    return container.metrics.snapshot()


_INGESTABLE_VISIBILITY = (Visibility.PRIVATE.value, Visibility.TENANT.value)


@router.post("/ingest", status_code=status.HTTP_202_ACCEPTED, dependencies=[Depends(admit_request)])
async def ingest(
    request: Request,
    file: UploadFile = File(...),
    scope: str | None = Form(None),
    visibility: str | None = Form(None),
    document_id: str | None = Form(None),
    principal: Principal = Depends(require_ingest),
    container: Container = Depends(get_container),
) -> dict:
    """Accept a file for ingestion, as a new document or a new version of one.

    Three outcomes, decided by `_resolve_target_document` + the content hash:

      - **created** -- nothing matched; a new document at version 1.
      - **updated** -- an existing document matched and the bytes differ; that
        document is repointed at the new bytes, `version` bumps, `updated_at` is
        stamped, and the job this returns re-runs the pipeline, which replaces
        its chunks and vectors. Superseded content stops being retrievable.
      - **deduplicated** -- the bytes are byte-for-byte what is already indexed.
        No job, no version bump; re-running the pipeline would be pure cost.

    `document_id` (optional) names the document this upload replaces. Omit it
    and the caller's own newest document with the same filename is used, which
    is what makes re-uploading `handbook.pdf` update it rather than create a
    twin. See `_resolve_target_document` for the authorization rules.

    `visibility` controls who inside the tenant may retrieve it:
      - `private` (default) -- only the uploader (and tenant admins).
      - `tenant`            -- everyone in the tenant.

    The default stays `private` so existing callers are unaffected and a
    personal upload is never exposed by accident; publishing to the whole tenant
    has to be asked for. `scope` is the orthogonal, cross-TENANT control and is
    unchanged. **On an update, both default to the document's CURRENT values
    rather than to `private`/`tenant`** -- silently re-privatising a
    tenant-shared document because someone re-uploaded it would be a regression
    disguised as a default.

    `scope`/`visibility` are read as `Optional[str] = Form(None)` rather than
    `str = Form(<default>)` on purpose: FastAPI/Starlette treats an explicitly
    submitted empty string as if the field were absent and silently substitutes
    the declared default, which let `visibility=""` sail through as `"private"`
    instead of being rejected like every other invalid value. Reading the raw
    value and applying our own default only when it is genuinely `None` (field
    omitted) keeps "omitted" and "sent empty" distinguishable -- and that same
    distinction is what lets an update inherit rather than reset.
    """
    cfg = container.settings
    form = await request.form()
    if "visibility" in form:
        visibility = form["visibility"]
    if "scope" in form:
        scope = form["scope"]
    bind(tenant_id=principal.tenant_id, user_id=principal.user_id)

    if visibility is not None and visibility not in _INGESTABLE_VISIBILITY:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            f"invalid visibility '{visibility}'. Expected one of {list(_INGESTABLE_VISIBILITY)}",
        )
    if scope is not None and scope not in (Scope.TENANT.value, Scope.GLOBAL.value):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"invalid scope '{scope}'")

    ext = os.path.splitext(file.filename or "")[1].lower()
    if ext not in EXT_TO_SOURCE_TYPE:
        raise HTTPException(
            status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            f"unsupported file type '{ext}'. Phase-1 accepts: {sorted(set(EXT_TO_SOURCE_TYPE))}",
        )
    data = await _read_capped(file, cfg.max_upload_mb * 1024 * 1024)
    if not data:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "empty file")

    return await run_in_threadpool(
        _accept_upload,
        container,
        principal,
        data,
        ext,
        file.filename or "",
        file.content_type or "application/octet-stream",
        document_id,
        visibility,
        scope,
    )


def _accept_upload(
    container: Container,
    principal: Principal,
    data: bytes,
    ext: str,
    filename: str,
    mime: str,
    document_id: str | None,
    visibility: str | None,
    scope: str | None,
) -> dict:
    cfg = container.settings

    source_type = EXT_TO_SOURCE_TYPE[ext].value
    sha256 = hashlib.sha256(data).hexdigest()

    container.metrics.incr("ingest.requests")

    target = _resolve_target_document(container, principal, document_id, filename)

    eff_visibility = visibility or (target.visibility if target else Visibility.PRIVATE.value)
    eff_scope = scope or (target.scope if target else Scope.TENANT.value)
    if eff_scope == Scope.GLOBAL.value and scope is not None:
        if not (
            cfg.platform_tenant_id
            and principal.tenant_id == cfg.platform_tenant_id
            and principal.role == Role.ADMIN
        ):
            raise HTTPException(
                status.HTTP_403_FORBIDDEN,
                "only the platform tenant's admin may publish to global scope",
            )

    if target is not None:
        if target.content_sha256 == sha256:
            if target.visibility != eff_visibility or target.scope != eff_scope:
                container.metadata.update_document_access(
                    principal.tenant_id, target.id, eff_visibility, eff_scope
                )

            container.metrics.incr("ingest.deduped")
            log.info(
                "ingest deduped",
                extra={"event": "ingest_deduped", "document_id": target.id, "sha": sha256[:12]},
            )
            return {
                "document_id": target.id,
                "deduplicated": True,
                "version": target.version,
                "updated_at": target.updated_at,
                "message": "identical content already ingested",
            }
        return _update_document(
            container,
            principal,
            target,
            data=data,
            ext=ext,
            sha256=sha256,
            source_type=source_type,
            filename=filename,
            mime=mime,
            visibility=eff_visibility,
            scope=eff_scope,
        )

    existing = container.metadata.get_document_by_hash(
        principal.tenant_id, sha256, principal.user_id
    )
    if existing is not None:
        container.metrics.incr("ingest.deduped")
        log.info(
            "ingest deduped",
            extra={"event": "ingest_deduped", "document_id": existing.id, "sha": sha256[:12]},
        )
        return {
            "document_id": existing.id,
            "deduplicated": True,
            "version": existing.version,
            "updated_at": existing.updated_at,
            "message": "identical content already ingested",
        }

    blob_path = container.blob.put(principal.tenant_id, sha256, ext, data)

    doc = Document(
        id=new_object_id(),
        tenant_id=principal.tenant_id,
        owner_user_id=principal.user_id,
        source_type=source_type,
        blob_path=blob_path,
        content_sha256=sha256,
        mime=mime,
        filename=filename,
        visibility=eff_visibility,
        acl_user_ids=[],
        scope=eff_scope,
    )
    job = Job(
        id=new_object_id(),
        document_id=doc.id,
        tenant_id=principal.tenant_id,
        stage=JobStage.PARSE.value,
        status=JobStatus.QUEUED.value,
        attempts=0,
    )

    initial_version = DocumentVersion(
        id=new_object_id(),
        document_id=doc.id,
        tenant_id=doc.tenant_id,
        version=1,
        content_sha256=doc.content_sha256,
        blob_path=doc.blob_path,
        filename=doc.filename,
        mime=doc.mime,
        byte_size=len(data),
        uploaded_by=principal.user_id,
        job_id=job.id,
    )
    container.metadata.create_document_with_job(doc, job, initial_version)
    container.metrics.incr("ingest.accepted")
    log.info(
        "ingest accepted",
        extra={
            "event": "ingest_accepted",
            "document_id": doc.id,
            "job_id": job.id,
            "source_type": source_type,
            "bytes": len(data),
            "sha": sha256[:12],
            "file": doc.filename,
        },
    )

    container.metadata.write_audit(
        principal.tenant_id,
        principal.user_id,
        "ingest",
        doc.id,
        {
            "source_type": source_type,
            "filename": doc.filename,
            "job_id": job.id,
            "scope": eff_scope,
            "visibility": eff_visibility,
        },
    )

    return {
        "job_id": job.id,
        "document_id": doc.id,
        "status": job.status,
        "source_type": source_type,
        "scope": eff_scope,
        "visibility": eff_visibility,
        "version": 1,
        "created": True,
    }


def _update_document(
    container: Container,
    principal: Principal,
    target: Document,
    *,
    data: bytes,
    ext: str,
    sha256: str,
    source_type: str,
    filename: str,
    mime: str,
    visibility: str,
    scope: str,
) -> dict:
    """Repoint an existing document at new bytes and queue the re-index.

    Only the `documents` row changes here. Its chunks and vectors still describe
    the OLD bytes until the worker runs the job this creates -- which is exactly
    what `_stage_upsert` handles, dropping every prior vector before writing the
    new set, so removed content stops being retrievable and new content starts.
    """
    if container.metadata.has_pending_job(principal.tenant_id, target.id):
        raise HTTPException(
            status.HTTP_409_CONFLICT,
            f"document {target.id} already has an ingestion job in progress; "
            f"wait for it to finish before uploading again",
        )

    other = container.metadata.get_document_by_hash(
        principal.tenant_id, sha256, target.owner_user_id
    )
    if other is not None and other.id != target.id:
        raise HTTPException(
            status.HTTP_409_CONFLICT,
            f"these bytes are already ingested as document {other.id}; "
            f"delete it first, or re-ingest under that document id",
        )

    previous_version = target.version
    previous_sha = target.content_sha256

    blob_path = container.blob.put(principal.tenant_id, sha256, ext, data)

    job = Job(
        id=new_object_id(),
        document_id=target.id,
        tenant_id=principal.tenant_id,
        stage=JobStage.PARSE.value,
        status=JobStatus.QUEUED.value,
        attempts=0,
    )

    version = container.metadata.update_document_content_and_queue(
        principal.tenant_id,
        target.id,
        blob_path=blob_path,
        content_sha256=sha256,
        mime=mime,
        source_type=source_type,
        filename=filename,
        visibility=visibility,
        scope=scope,
        job=job,
        uploaded_by=principal.user_id,
        byte_size=len(data),
        expected_version=target.version,
    )

    container.metrics.incr("ingest.accepted")
    container.metrics.incr("ingest.updated")
    log.info(
        "ingest update accepted",
        extra={
            "event": "ingest_updated",
            "document_id": target.id,
            "job_id": job.id,
            "source_type": source_type,
            "bytes": len(data),
            "sha": sha256[:12],
            "file": filename,
            "version": version,
        },
    )
    container.metadata.write_audit(
        principal.tenant_id,
        principal.user_id,
        "update",
        target.id,
        {
            "source_type": source_type,
            "filename": filename,
            "job_id": job.id,
            "scope": scope,
            "visibility": visibility,
            "version": version,
            "previous_version": previous_version,
            "previous_sha": previous_sha[:12],
        },
    )
    return {
        "job_id": job.id,
        "document_id": target.id,
        "status": job.status,
        "source_type": source_type,
        "scope": scope,
        "visibility": visibility,
        "version": version,
        "previous_version": previous_version,
        "updated": True,
    }


@router.get("/jobs/{job_id}")
def get_job(
    job_id: str,
    principal: Principal = Depends(get_principal),
    container: Container = Depends(get_container),
) -> dict:
    """Status of one ingestion job.

    Gated on the DOCUMENT's visibility, not just on tenant membership: the
    response carries the document id, its pipeline stage and any error text, so
    a job is exactly as sensitive as the document behind it. `/jobs/{id}/trace`
    already enforced this; this endpoint did not, which let any member of a
    tenant watch another user's private ingest.
    """
    job = container.metadata.get_job(principal.tenant_id, job_id)
    if job is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "job not found")
    doc = container.metadata.get_document(principal.tenant_id, job.document_id)
    if doc is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "job not found")
    _visible_or_404(doc, principal)
    return {
        "job_id": job.id,
        "document_id": job.document_id,
        "stage": job.stage,
        "status": job.status,
        "attempts": job.attempts,
        "error": job.error,
        "route_summary": job.route_summary,
        "created_at": job.created_at,
        "updated_at": job.updated_at,
    }


@router.get("/jobs/{job_id}/trace")
def get_job_trace(
    job_id: str,
    principal: Principal = Depends(get_principal),
    container: Container = Depends(get_container),
) -> dict:
    """Everything the trace view needs for one ingestion, in a single call:
    where the job is, how long each stage took and what it produced, how the
    document was routed, and what ended up indexed."""
    job = container.metadata.get_job(principal.tenant_id, job_id)
    if job is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "job not found")
    doc = container.metadata.get_document(principal.tenant_id, job.document_id)
    if doc is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "document not found")
    _visible_or_404(doc, principal)

    events = container.metadata.get_job_events(principal.tenant_id, job_id)
    chunks = container.metadata.get_document_chunks(principal.tenant_id, doc.id)
    return {
        "job": {
            "job_id": job.id,
            "document_id": job.document_id,
            "stage": job.stage,
            "status": job.status,
            "attempts": job.attempts,
            "error": job.error,
            "route_summary": job.route_summary,
            "created_at": job.created_at,
            "updated_at": job.updated_at,
        },
        "document": {
            "document_id": doc.id,
            "filename": doc.filename,
            "source_type": doc.source_type,
            "mime": doc.mime,
            "scope": doc.scope,
            "visibility": doc.visibility,
            "extracted_metadata": doc.extracted_metadata,
            "created_at": doc.created_at,
            "version": doc.version,
            "updated_at": doc.updated_at,
        },
        "stages": [
            {
                "stage": e.stage,
                "seq": e.seq,
                "status": e.status,
                "duration_ms": e.duration_ms,
                "attempt": e.attempt,
                "detail": e.detail,
                "at": e.at,
            }
            for e in events
        ],
        "chunk_count": len(chunks),
        "token_total": sum(c.token_count or 0 for c in chunks),
    }


@router.get("/documents")
def list_documents(
    limit: int = 50,
    offset: int = 0,
    principal: Principal = Depends(get_principal),
    container: Container = Depends(get_container),
) -> dict:
    """Documents this caller may read, newest first, each with its latest job.
    Backs the trace view's history list.

    Note: the store pages by tenant, then `can_view` filters the page, so a page
    can come back shorter than `limit` without meaning the end was reached. That
    matches how retrieval already filters post-fetch (`access_predicate`), and
    keeps the ACL rule in exactly one place. Proper ACL-aware pagination is part
    of the "response pagination on list endpoints" roadmap item."""
    limit = max(1, min(limit, 200))
    rows = container.metadata.list_documents(principal.tenant_id, limit, max(0, offset))
    items = []
    for doc, job in rows:
        payload = {
            "user_id": doc.owner_user_id,
            "visibility": doc.visibility,
            "acl_user_ids": doc.acl_user_ids,
            "scope": doc.scope,
        }
        if not can_view(payload, principal.user_id, principal.role.value):
            continue
        items.append(
            {
                "document_id": doc.id,
                "filename": doc.filename,
                "source_type": doc.source_type,
                "scope": doc.scope,
                "visibility": doc.visibility,
                "created_at": doc.created_at,
                "version": doc.version,
                "updated_at": doc.updated_at,
                "active_generation_id": doc.active_generation_id,
                "indexed_version": doc.indexed_version,
                "job_id": job.id if job else None,
                "status": job.status if job else None,
                "stage": job.stage if job else None,
            }
        )
    return {"documents": items, "count": len(items), "limit": limit, "offset": max(0, offset)}


@router.get("/documents/{document_id}")
def get_document(
    document_id: str,
    principal: Principal = Depends(get_principal),
    container: Container = Depends(get_container),
) -> dict:
    """One document's metadata plus the tenant's total indexed-vector count."""
    doc = container.metadata.get_document(principal.tenant_id, document_id)
    if doc is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "document not found")
    _visible_or_404(doc, principal)
    return {
        "document_id": doc.id,
        "source_type": doc.source_type,
        "filename": doc.filename,
        "visibility": doc.visibility,
        "scope": doc.scope,
        "mime": doc.mime,
        "extracted_metadata": doc.extracted_metadata,
        "vector_count": container.vectors.count(principal.tenant_id),
        "created_at": doc.created_at,
        "version": doc.version,
        "updated_at": doc.updated_at,
    }


@router.get("/documents/{document_id}/versions")
def list_document_versions(
    document_id: str,
    principal: Principal = Depends(get_principal),
    container: Container = Depends(get_container),
) -> dict:
    """A document's content-version history, newest first.

    Each entry is a set of bytes that was ingested under this document id, and
    how the chunk set changed as a result. Historical source bytes remain
    viewable, while only the newest version's chunks and vectors are searchable.

    The `chunks_*` fields are null on a version whose job hasn't finished yet
    (or failed): the delta is written by the worker after the chunk stage.
    """
    doc = container.metadata.get_document(principal.tenant_id, document_id)
    if doc is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "document not found")
    _visible_or_404(doc, principal)
    versions = container.metadata.list_document_versions(doc.tenant_id, document_id)
    return {
        "document_id": document_id,
        "version": doc.version,
        "updated_at": doc.updated_at,
        "count": len(versions),
        "versions": [
            {
                "version": v.version,
                "sha": v.content_sha256[:12],
                "filename": v.filename,
                "mime": v.mime,
                "byte_size": v.byte_size,
                "uploaded_by": v.uploaded_by,
                "job_id": v.job_id,
                "created_at": v.created_at,
                "current": v.content_sha256 == doc.content_sha256,
                "chunks_added": v.chunks_added,
                "chunks_removed": v.chunks_removed,
                "chunks_unchanged": v.chunks_unchanged,
                "added_preview": (v.delta or {}).get("added_preview") or [],
                "removed_preview": (v.delta or {}).get("removed_preview") or [],
                "added_truncated": bool((v.delta or {}).get("added_truncated")),
                "removed_truncated": bool((v.delta or {}).get("removed_truncated")),
            }
            for v in versions
        ],
    }


def _source_version(
    container: Container, principal: Principal, document_id: str, version: int | None
):
    doc = container.metadata.get_document(principal.tenant_id, document_id)
    if doc is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "document not found")
    _visible_or_404(doc, principal)
    selected = version or doc.version
    item = container.metadata.get_document_version(doc.tenant_id, doc.id, selected)
    if item is None and selected == doc.version:
        item = DocumentVersion(
            id=f"{doc.id}:v{doc.version}",
            document_id=doc.id,
            tenant_id=doc.tenant_id,
            version=doc.version,
            content_sha256=doc.content_sha256,
            blob_path=doc.blob_path,
            filename=doc.filename,
            byte_size=0,
            uploaded_by=doc.owner_user_id,
            mime=doc.mime,
        )
    if item is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "document version not found")
    return doc, item


def _byte_range(value: str | None, size: int) -> tuple[int, int] | None:
    if not value:
        return None
    match = re.fullmatch(r"bytes=(\d*)-(\d*)", value.strip())
    if not match or size <= 0:
        raise ValueError
    first, last = match.groups()
    if not first:
        length = int(last or 0)
        if length <= 0:
            raise ValueError
        return max(0, size - length), size - 1
    start = int(first)
    end = min(int(last), size - 1) if last else size - 1
    if start >= size or end < start:
        raise ValueError
    return start, end


def _source_response(
    container: Container, item: DocumentVersion, range_header: str | None
) -> Response:
    try:
        size = container.blob.size(item.blob_path)
    except (FileNotFoundError, ValueError, OSError) as exc:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "source content not found") from exc
    headers = {
        "Accept-Ranges": "bytes",
        "Content-Disposition": f"inline; filename*=UTF-8''{quote(item.filename)}",
        "Cache-Control": "private, no-store",
        "ETag": f'"{item.content_sha256}"',
        "X-Content-Type-Options": "nosniff",
    }
    try:
        selected = _byte_range(range_header, size)
    except ValueError as exc:
        raise HTTPException(
            status.HTTP_416_RANGE_NOT_SATISFIABLE,
            "invalid byte range",
            headers={"Content-Range": f"bytes */{size}"},
        ) from exc
    if selected is None:
        try:
            data = container.blob.get(item.blob_path)
        except (FileNotFoundError, ValueError, OSError) as exc:
            raise HTTPException(status.HTTP_404_NOT_FOUND, "source content not found") from exc
        headers["Content-Length"] = str(len(data))
        media_type = (
            item.mime
            if item.mime != "application/octet-stream"
            else (mimetypes.guess_type(item.filename)[0] or item.mime)
        )
        return Response(data, media_type=media_type, headers=headers)
    start, end = selected
    try:
        data = container.blob.get_range(item.blob_path, start, end)
    except (FileNotFoundError, ValueError, OSError) as exc:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "source content not found") from exc
    headers.update(
        {"Content-Range": f"bytes {start}-{end}/{size}", "Content-Length": str(len(data))}
    )
    media_type = (
        item.mime
        if item.mime != "application/octet-stream"
        else (mimetypes.guess_type(item.filename)[0] or item.mime)
    )
    return Response(data, status_code=206, media_type=media_type, headers=headers)


@router.get("/documents/{document_id}/content")
def get_current_document_content(
    document_id: str,
    request: Request,
    principal: Principal = Depends(get_principal),
    container: Container = Depends(get_container),
):
    _, item = _source_version(container, principal, document_id, None)
    return _source_response(container, item, request.headers.get("range"))


@router.get("/documents/{document_id}/versions/{version}/content")
def get_document_version_content(
    document_id: str,
    version: int,
    request: Request,
    principal: Principal = Depends(get_principal),
    container: Container = Depends(get_container),
):
    _, item = _source_version(container, principal, document_id, version)
    return _source_response(container, item, request.headers.get("range"))


@router.get("/documents/{document_id}/versions/{version}/pages/{page_number}")
def render_pdf_page(
    document_id: str,
    version: int,
    page_number: int,
    principal: Principal = Depends(get_principal),
    container: Container = Depends(get_container),
):
    _, item = _source_version(container, principal, document_id, version)
    if not item.filename.lower().endswith(".pdf"):
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST, "page rendering is available for PDF documents"
        )

    try:
        data = container.blob.get(item.blob_path)
    except (FileNotFoundError, ValueError, OSError) as exc:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "source content not found") from exc
    document = pdfium.PdfDocument(data)
    try:
        if page_number < 1 or page_number > len(document):
            raise HTTPException(status.HTTP_404_NOT_FOUND, "page not found")
        bitmap = document[page_number - 1].render(scale=2)
        image = bitmap.to_pil()
        output = io.BytesIO()
        image.save(output, "PNG")
    finally:
        document.close()
    return Response(
        output.getvalue(),
        media_type="image/png",
        headers={"Cache-Control": "private, no-store", "X-Content-Type-Options": "nosniff"},
    )


@router.get("/documents/{document_id}/versions/{version}/preview")
def preview_document_version(
    document_id: str,
    version: int,
    principal: Principal = Depends(get_principal),
    container: Container = Depends(get_container),
):
    """Deterministic text preview of the exact stored version for non-native formats."""
    _, item = _source_version(container, principal, document_id, version)
    try:
        data = container.blob.get(item.blob_path)
    except (FileNotFoundError, ValueError, OSError) as exc:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "source content not found") from exc
    ext = os.path.splitext(item.filename)[1].lower()
    source_blocks = []
    try:
        if ext == ".docx":
            document = Docx(io.BytesIO(data))
            blocks = []
            for body_index, child in enumerate(document.element.body.iterchildren(), 1):
                if child.tag.endswith("}p"):
                    value = Paragraph(child, document).text.strip()
                elif child.tag.endswith("}tbl"):
                    value = "\n".join(
                        "\t".join(cell.text for cell in row.cells)
                        for row in Table(child, document).rows
                    )
                else:
                    value = ""
                if value:
                    blocks.append(value)
                    source_blocks.append({"text": value, "locator": {"body_index": body_index}})
            text = "\n\n".join(blocks)
        elif ext == ".rtf":
            text = rtf_to_text(data.decode("utf-8-sig"))
        elif ext == ".xlsx":
            workbook = load_workbook(io.BytesIO(data), data_only=False, read_only=True)
            try:
                sections = []
                for sheet in workbook.worksheets:
                    if (
                        sheet.max_row > container.settings.max_table_rows
                        or sheet.max_row * sheet.max_column > container.settings.max_workbook_cells
                    ):
                        raise ValueError("Workbook preview exceeds cell limits")
                    rows = [
                        "\t".join("" if value is None else str(value) for value in row)
                        for row in sheet.iter_rows(values_only=True)
                    ]
                    sections.append(f"[{sheet.title}]\n" + "\n".join(rows))
                    source_blocks.extend(
                        {"text": row, "locator": {"sheet": sheet.title, "row": index}}
                        for index, row in enumerate(rows, 1)
                    )
                text = "\n\n".join(sections)
            finally:
                workbook.close()
        elif ext == ".xls":
            workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
            try:
                sections = []
                for sheet in workbook.sheets():
                    if (
                        sheet.nrows > container.settings.max_table_rows
                        or sheet.nrows * sheet.ncols > container.settings.max_workbook_cells
                    ):
                        raise ValueError("Workbook preview exceeds cell limits")
                    rows = [
                        "\t".join(
                            str(sheet.cell_value(row, column)) for column in range(sheet.ncols)
                        )
                        for row in range(sheet.nrows)
                    ]
                    sections.append(f"[{sheet.name}]\n" + "\n".join(rows))
                    source_blocks.extend(
                        {"text": row, "locator": {"sheet": sheet.name, "row": index}}
                        for index, row in enumerate(rows, 1)
                    )
                text = "\n\n".join(sections)
            finally:
                workbook.release_resources()
        elif ext == ".csv":
            rows = []
            for index, row in enumerate(csv.reader(io.StringIO(data.decode("utf-8-sig"))), 1):
                if index > container.settings.max_table_rows + 1:
                    raise ValueError("CSV preview exceeds row limit")
                value = "\t".join(row)
                rows.append(value)
                source_blocks.append({"text": value, "locator": {"row": index}})
            text = "\n".join(rows)
        else:
            text = data.decode("utf-8-sig")
    except Exception as exc:
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_ENTITY, "source preview could not be generated"
        ) from exc
    limit = 2_000_000
    bounded_blocks, used = [], 0
    for block in source_blocks:
        if used >= limit or len(bounded_blocks) >= 2000:
            break
        value = block["text"][: limit - used]
        bounded_blocks.append({**block, "text": value})
        used += len(value)
    return {
        "text": text[:limit],
        "truncated": len(text) > limit,
        "filename": item.filename,
        "version": item.version,
        "blocks": bounded_blocks,
        "blocks_truncated": len(bounded_blocks) < len(source_blocks),
    }


@router.get("/documents/{document_id}/versions/{version}/frames/{frame_number}")
def render_image_frame(
    document_id: str,
    version: int,
    frame_number: int,
    principal: Principal = Depends(get_principal),
    container: Container = Depends(get_container),
):
    """Render the cited image frame as PNG, including browser-incompatible TIFF."""
    _, item = _source_version(container, principal, document_id, version)
    if os.path.splitext(item.filename)[1].lower() not in {
        ".png",
        ".jpg",
        ".jpeg",
        ".webp",
        ".tif",
        ".tiff",
    }:
        raise HTTPException(400, "frame rendering requires an image document")

    try:
        data = container.blob.get(item.blob_path)
        with Image.open(io.BytesIO(data)) as image:
            if not 1 <= frame_number <= getattr(image, "n_frames", 1):
                raise HTTPException(404, "frame not found")
            image.seek(frame_number - 1)
            if image.width * image.height > container.settings.max_image_pixels:
                raise HTTPException(413, "image frame exceeds pixel limit")
            with ImageOps.exif_transpose(image).convert("RGB") as frame:
                output = io.BytesIO()
                frame.save(output, "PNG")
    except (OSError, ValueError, Image.DecompressionBombError) as exc:
        raise HTTPException(422, "image preview unavailable") from exc
    return Response(
        output.getvalue(),
        media_type="image/png",
        headers={"Cache-Control": "private, no-store", "X-Content-Type-Options": "nosniff"},
    )


@router.post("/documents/{document_id}/reprocess", status_code=status.HTTP_202_ACCEPTED)
def reprocess(
    document_id: str,
    principal: Principal = Depends(require_ingest),
    container: Container = Depends(get_container),
) -> dict:
    """Re-run the pipeline on an existing document (e.g. after a chunking change).
    Upsert drops prior vectors first, so it's idempotent.

    Requires BOTH checks: `_visible_or_404` (you may not act on a document you
    are not allowed to see -- tenant membership alone previously let any member
    reprocess another user's private document) and `_owned_or_404` (read access
    to a global document does not grant the right to reprocess it).
    """
    doc = container.metadata.get_document(principal.tenant_id, document_id)
    if doc is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "document not found")
    _visible_or_404(doc, principal)
    _owned_or_404(doc, principal)
    if container.metadata.has_pending_job(principal.tenant_id, doc.id):
        raise HTTPException(
            status.HTTP_409_CONFLICT,
            f"document {doc.id} already has an ingestion job in progress; "
            f"wait for it to finish before reprocessing",
        )

    job = Job(
        id=new_object_id(),
        document_id=doc.id,
        tenant_id=principal.tenant_id,
        stage=JobStage.PARSE.value,
        status=JobStatus.QUEUED.value,
        attempts=0,
        blob_path=doc.blob_path,
        content_sha256=doc.content_sha256,
        version=doc.version,
    )
    container.metadata.create_job(job)
    container.metrics.incr("reprocess.requests")
    log.info(
        "reprocess queued", extra={"event": "reprocess", "document_id": doc.id, "job_id": job.id}
    )
    return {"job_id": job.id, "document_id": doc.id, "status": job.status}


@router.get("/documents/{document_id}/chunks")
def list_chunks(
    document_id: str,
    preview: int = 240,
    principal: Principal = Depends(get_principal),
    container: Container = Depends(get_container),
) -> dict:
    """Chunk-level breakdown of one document, with a truncated text preview per chunk."""
    doc = container.metadata.get_document(principal.tenant_id, document_id)
    if doc is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "document not found")
    _visible_or_404(doc, principal)
    recs = container.metadata.get_document_chunks(doc.tenant_id, document_id)
    return {
        "document_id": document_id,
        "chunk_count": len(recs),
        "chunks": [
            {
                "chunk_id": r.id,
                "content_sha256": r.content_sha256,
                "ordinal": r.ordinal,
                "modality": r.modality,
                "extractor": r.extractor,
                "route_reason": r.route_reason,
                "tokens": r.token_count,
                "pages": r.meta.get("pages") or r.meta.get("page"),
                "location": location_str(r.meta),
                "provenance": citation_provenance(r.meta, doc.source_type),
                "preview": r.text[:preview],
            }
            for r in recs
        ],
    }


@router.delete("/documents/{document_id}")
def delete_document(
    document_id: str,
    principal: Principal = Depends(require_delete),
    container: Container = Depends(get_container),
) -> dict:
    """Delete a document, its vectors, every version's blob, and invalidate any
    cached answers it fed."""
    doc = container.metadata.get_document(principal.tenant_id, document_id)
    if doc is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "document not found")
    _owned_or_404(doc, principal)
    removed = container.vectors.delete_by_document(principal.tenant_id, document_id)

    versions = container.metadata.list_document_versions(doc.tenant_id, document_id)
    blob_paths = {doc.blob_path} | {v.blob_path for v in versions if v.blob_path}
    container.metadata.delete_document(principal.tenant_id, document_id)
    for path in blob_paths:
        _delete_blob_if_unreferenced(container, path, document_id)

    try:
        invalidate_cache_for(container, principal.tenant_id, doc.scope)
    except Exception:
        log.warning(
            "Deleted document cache invalidation failed; durable epoch remains authoritative",
            extra={
                "event": "delete_cache_invalidation_failed",
                "document_id": document_id,
                "tenant_id": principal.tenant_id,
            },
        )
    container.metrics.incr("documents.deleted")
    container.metadata.write_audit(
        principal.tenant_id,
        principal.user_id,
        "delete",
        document_id,
        {"vectors_removed": removed},
    )
    return {"document_id": document_id, "deleted": True, "vectors_removed": removed}


def _delete_blob_if_unreferenced(container: Container, blob_path: str, document_id: str) -> None:
    """Unlink a blob only if no other document or version still points at it.

    Blobs are content-addressed per tenant, so identical bytes are one file on
    disk no matter how many documents reference them. `UNIQUE (tenant_id,
    content_sha256)` on `documents` used to make sharing impossible, which is
    why this was an unconditional unlink -- but versioning breaks that
    assumption: once v1's bytes are no longer any document's CURRENT hash, the
    ingest dedup check stops matching them, so the very same bytes can come back
    as a brand-new document while this document's v1 history row still names the
    file. Unlinking it then would silently empty out the other document.

    Best-effort, like the rest of the delete path: the metadata rows are already
    gone by this point, so a failure here leaks a file rather than corrupting
    state, and must not turn a successful delete into a 500.
    """
    try:
        if container.metadata.count_blob_references(blob_path, document_id) == 0:
            container.blob.delete(blob_path)
    except Exception:  # noqa: BLE001
        log.warning(
            "blob not deleted",
            extra={"event": "blob_delete_failed", "document_id": document_id},
            exc_info=True,
        )
