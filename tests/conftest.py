"""Shared test fixtures.

Every test runs against a hermetic container rooted at a per-test tmp dir, with
the deterministic local embedder (no gateway needed). File fixtures are built in
memory so tests carry no binary blobs.
"""
from __future__ import annotations

import io

import pytest

from app.config import Settings
from app.container import build_container
from app.domain.models import Role


@pytest.fixture
def settings(tmp_path):
    return Settings(
        data_dir=tmp_path / "data",
        metadata_backend="sqlite",
        blob_backend="localfs",
        vector_backend="localfile",
        queue_backend="sqlite",
        embedding_provider="minilm",
        litellm_api_key="",
        vision_model="test-vision",
        max_attempts=3,
        worker_poll_seconds=0.01,
        # reranking is opt-in per-test (see tests/test_query_rerank.py) so the
        # rest of the suite isn't coupled to a second real model download/load.
        reranker_provider="none",
    )


@pytest.fixture
def container(settings):
    # real MiniLM embedder (ONNX). The model is process-cached, so it loads once
    # for the whole test session.
    c = build_container(settings)
    # Default chat() stub so every test stays offline (no real gateway call) — the
    # metadata-extraction pipeline stage calls chat() on every ingest. Tests that
    # care about a specific chat response (query answers, populated metadata)
    # override this via monkeypatch, same as before this stub existed.
    c.gateway.chat = lambda messages, model, temperature=0.0: (
        '{"author": null, "date": null, "topics": [], "entities": []}'
    )
    return c


@pytest.fixture
def tenant(container):
    """A tenant with admin/member/viewer users; returns ids + tokens."""
    tid = container.metadata.create_tenant("Acme")
    admin = "sk-admin"
    member = "sk-member"
    viewer = "sk-viewer"
    admin_id = container.metadata.create_user(tid, "admin@acme.test", Role.ADMIN.value, admin)
    member_id = container.metadata.create_user(tid, "member@acme.test", Role.MEMBER.value, member)
    viewer_id = container.metadata.create_user(tid, "viewer@acme.test", Role.VIEWER.value, viewer)
    return {
        "id": tid,
        "admin_token": admin, "admin_id": admin_id,
        "member_token": member, "member_id": member_id,
        "viewer_token": viewer, "viewer_id": viewer_id,
    }


@pytest.fixture
def other_tenant(container):
    tid = container.metadata.create_tenant("Globex")
    token = "sk-other-admin"
    container.metadata.create_user(tid, "admin@globex.test", Role.ADMIN.value, token)
    return {"id": tid, "admin_token": token}


def _csv_bytes() -> bytes:
    rows = "\n".join(f"{i},item{i},{i * 10}" for i in range(1, 21))
    return ("id,name,value\n" + rows + "\n").encode()


def _txt_bytes() -> bytes:
    para = ("The quarterly review covers revenue and costs. "
            "Revenue grew across regions with APAC leading. ") * 4
    return ("\n\n".join([para, para, "# Outlook\n\n" + para])).encode()


def _docx_bytes() -> bytes:
    from docx import Document
    d = Document()
    d.add_heading("Annual Report", level=1)
    d.add_paragraph("Revenue grew across all regions this year. " * 6)
    d.add_heading("Financials", level=2)
    d.add_paragraph("Costs were controlled while margins improved. " * 6)
    t = d.add_table(rows=1, cols=3)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text, t.rows[0].cells[2].text = \
        "Region", "Revenue", "Growth"
    for r in [("APAC", "120", "18%"), ("EMEA", "90", "9%")]:
        c = t.add_row().cells
        c[0].text, c[1].text, c[2].text = r
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xlsx_bytes() -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Q1", "Q2"])
    for r in [("APAC", 100, 120), ("EMEA", 80, 90)]:
        ws.append(r)
    ws2 = wb.create_sheet("Costs")
    ws2.append(["Dept", "Budget"])
    ws2.append(["R&D", 50])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pdf_bytes() -> bytes:
    from fpdf import FPDF
    para = ("The quarterly review covers revenue and costs. "
            "Revenue grew across regions. ") * 4
    pdf = FPDF()
    pdf.set_auto_page_break(True, 15)
    pdf.set_font("Helvetica", size=12)
    for pg in range(2):
        pdf.add_page()
        pdf.multi_cell(0, 8, f"Page {pg + 1}\n\n{para}")
    return bytes(pdf.output())


@pytest.fixture
def files() -> dict[str, bytes]:
    return {
        "data.csv": _csv_bytes(),
        "notes.txt": _txt_bytes(),
        "report.docx": _docx_bytes(),
        "finance.xlsx": _xlsx_bytes(),
        "doc.pdf": _pdf_bytes(),
    }


class FakeGateway:
    """Stand-in for the LiteLLM client in tests that exercise vision routing."""
    def __init__(self, vision_return: str = "A bar chart of quarterly revenue."):
        self._v = vision_return
        self.vision_calls = 0

    def vision(self, image_bytes: bytes, prompt: str, mime: str = "image/png") -> str:
        self.vision_calls += 1
        return self._v

    def ocr(self, image_bytes: bytes, mime: str = "image/png") -> str:
        return self.vision(image_bytes, "", mime)

    def embed(self, texts: list[str]):
        raise AssertionError("embed should go through the Embedder, not the gateway")


class NoGateway:
    """Fails loudly if any gateway call happens (offline-path assertion)."""
    def vision(self, *a, **k):
        raise AssertionError("gateway.vision called on an offline file")

    def embed(self, *a, **k):
        raise AssertionError("gateway.embed called")
