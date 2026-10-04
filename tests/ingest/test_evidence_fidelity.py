import csv
import hashlib
import io
import zipfile
from datetime import datetime

import pytest
from docx import Document as DocxDocument
from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.chart import BarChart
from PIL import Image

from app.ingest.pipeline.elements import Element
from app.ingest.pipeline.loaders import extract_document
from app.ingest.pipeline.loaders.excel import _display_value
from app.ingest.pipeline.loaders.pdf import _reading_order
from app.ingest.pipeline.runner import _stage_chunk, run_job
from app.ingest.pipeline.safety import UnsafeContentError
from app.shared.config import Settings
from app.shared.domain.models import Document, Job
from app.shared.ids import new_object_id
from tests.conftest import FakeGateway, NoGateway

CFG = Settings(_env_file=None)


@pytest.mark.parametrize(
    "value,pattern,expected",
    [
        (1234.565, "$#,##0.00", "$1,234.57"),
        (0.075, "0.00%", "7.50%"),
        (7, "000", "007"),
    ],
)
def test_business_numeric_display_formats(value, pattern, expected):
    assert _display_value(value, pattern) == expected


def test_business_date_display_is_distinct_from_typed_date():

    assert _display_value(datetime(2026, 9, 30), "mm/dd/yyyy") == "09/30/2026"


def test_csv_preserves_identifiers_na_and_decimal_strings():
    elements, _ = extract_document(
        "evidence.csv", b"id,status,amount\n00123,NA,007\n00456,N/A,12.00\n", NoGateway(), CFG
    )
    assert "| 00123 | NA | 007 |" in elements[0].text
    assert "| 00456 | N/A | 12.00 |" in elements[0].text


def test_html_preserves_prose_and_table_order():
    source = b"<html><body><h1>Invoices</h1><p>Account <b>00123</b> due</p><table><tr><th>ID</th><th>Total</th></tr><tr><td>007</td><td>12.00</td></tr></table><p>Pay by 2026-09-30.</p></body></html>"
    elements, _ = extract_document("evidence.html", source, NoGateway(), CFG)
    assert [e.text for e in elements] == [
        "# Invoices",
        "Account 00123 due",
        "| ID | Total |\n| --- | --- |\n| 007 | 12.00 |",
        "Pay by 2026-09-30.",
    ]


def test_workbook_separates_display_typed_formula_values():

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["ID", "Amount", "Formula"])
    sheet.append([123, 7, "=B2*2"])
    sheet["A2"].number_format = "00000"
    sheet["B2"].number_format = "0.00"
    buffer = io.BytesIO()
    workbook.save(buffer)
    elements, _ = extract_document("evidence.xlsx", buffer.getvalue(), NoGateway(), CFG)
    assert "00123" in elements[0].text
    assert "7.00" in elements[0].text
    cells = {c["address"]: c for c in elements[0].meta["cells"]}
    assert cells["A2"]["typed_value"] == 123
    assert cells["A2"]["display_value"] == "00123"
    assert cells["C2"]["formula"] == "=B2*2"
    assert cells["C2"]["typed_value"] is None
    assert "cached value unavailable" in cells["C2"]["display_value"]


def test_docx_figure_retains_anchor_section_and_order():

    document = DocxDocument()
    document.add_heading("Invoice detail", 1)
    document.add_paragraph("Before figure 00123")
    picture = io.BytesIO()
    Image.new("RGB", (16, 16), "white").save(picture, format="PNG")
    picture.seek(0)
    document.add_picture(picture)
    document.add_paragraph("After figure 007")
    buffer = io.BytesIO()
    document.save(buffer)
    elements, _ = extract_document(
        "evidence.docx", buffer.getvalue(), FakeGateway("Total 12.00 USD"), CFG
    )
    assert [e.text for e in elements] == [
        "# Invoice detail",
        "Before figure 00123",
        "Total 12.00 USD",
        "After figure 007",
    ]
    assert elements[2].meta["section_path"] == "# Invoice detail"
    assert [e.order for e in elements] == list(range(len(elements)))


def test_tiff_frames_are_normalized_and_preserve_order():
    buffer = io.BytesIO()
    Image.new("RGB", (16, 16), "white").save(
        buffer, format="TIFF", save_all=True, append_images=[Image.new("RGB", (16, 16), "black")]
    )

    class Gateway:
        calls = 0

        def vision(self, image_bytes, prompt, mime):
            assert mime == "image/png"
            assert image_bytes.startswith(b"\x89PNG")
            self.calls += 1
            return f"Frame {self.calls} ID 00{self.calls}"

    elements, _ = extract_document("evidence.tiff", buffer.getvalue(), Gateway(), CFG)
    assert [e.text for e in elements] == ["Frame 1 ID 001", "Frame 2 ID 002"]
    assert [e.meta["frame"] for e in elements] == [1, 2]


def test_native_pdf_two_columns_do_not_interleave():

    document = FPDF()
    document.add_page()
    document.set_font("Helvetica", size=12)
    for x, y, text in [
        (15, 20, "Left ID 00123"),
        (15, 30, "Left amount 007.00"),
        (120, 20, "Right ID 00456"),
        (120, 30, "Right amount 008.00"),
    ]:
        document.text(x, y, text)
    elements, _ = extract_document("columns.pdf", bytes(document.output()), NoGateway(), CFG)
    assert [e.text for e in elements] == [
        "Left ID 00123",
        "Left amount 007.00",
        "Right ID 00456",
        "Right amount 008.00",
    ]
    assert all(e.meta["page"] == 1 for e in elements)
    assert elements[0].meta["bbox"][2] < elements[2].meta["bbox"][0]


def test_pdf_full_width_heading_precedes_column_regions():

    boxes = [
        ("Left 001", (10, 30, 80, 40)),
        ("Right 003", (120, 30, 190, 40)),
        ("Title", (10, 10, 190, 20)),
        ("Right 004", (120, 50, 190, 60)),
        ("Left 002", (10, 50, 80, 60)),
    ]
    regions = [
        ((box[1], box[0]), Element(text, "text", "pdf_text", "text_layer", 0, {"bbox": box}))
        for text, box in boxes
    ]
    assert [element.text for _, element in _reading_order(regions, 200)] == [
        "Title",
        "Left 001",
        "Left 002",
        "Right 003",
        "Right 004",
    ]


@pytest.mark.parametrize("filename", ["evidence.txt", "evidence.md"])
def test_invalid_encoding_is_never_silently_dropped(filename):
    with pytest.raises(UnicodeDecodeError):
        extract_document(filename, b"Account 00123 \xff 007", NoGateway(), CFG)


def test_docx_bad_image_does_not_hide_later_image_or_prose():

    document = DocxDocument()
    paragraph = document.add_paragraph("Before 00123 ")
    for color in ("white", "black"):
        picture = io.BytesIO()
        Image.new("RGB", (16, 16), color).save(picture, format="PNG")
        picture.seek(0)
        paragraph.add_run().add_picture(picture)
    paragraph.add_run(" After 007")
    original = io.BytesIO()
    document.save(original)
    modified = io.BytesIO()
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(modified, "w") as target:
        for info in source.infolist():
            target.writestr(
                info,
                b"broken image" if info.filename == "word/media/image1.png" else source.read(info),
            )
    gateway = FakeGateway("Second image total 12.00")
    elements, summary = extract_document("figures.docx", modified.getvalue(), gateway, CFG)
    assert [e.text for e in elements if e.text] == [
        "Before 00123",
        "Second image total 12.00",
        "After 007",
    ]
    assert summary["omitted_regions"] == 1
    assert gateway.vision_calls == 1


def test_omission_only_source_cannot_be_indexed_empty():

    book = Workbook()
    book.active.add_chart(BarChart(), "A1")
    buffer = io.BytesIO()
    book.save(buffer)
    with pytest.raises(UnsafeContentError, match="no extracted evidence"):
        extract_document("chart.xlsx", buffer.getvalue(), NoGateway(), CFG)


@pytest.mark.parametrize(
    "filename,source,expected",
    [
        ("ledger.csv", b"id,status,amount\n00123,NA,007.00\n", ("00123", "NA", "007.00")),
        (
            "ledger.html",
            b"<h1>Ledger</h1><p>Invoice 00123 totals 007.00 USD on 2026-09-30.</p>",
            ("00123", "007.00", "2026-09-30"),
        ),
    ],
)
def test_exact_business_evidence_survives_parse_store_retrieve(
    container, tenant, filename, source, expected
):

    digest = hashlib.sha256(source).hexdigest()
    blob = container.blob.put(tenant["id"], digest, "." + filename.rsplit(".", 1)[1], source)
    document = Document(
        new_object_id(),
        tenant["id"],
        tenant["admin_id"],
        "table",
        blob,
        digest,
        "text/plain",
        filename,
        "tenant",
        [],
    )
    job = Job(new_object_id(), document.id, tenant["id"], "parse", "queued", 0)
    container.metadata.create_document_with_job(document, job)
    run_job(container, container.queue.claim_next())
    records = container.metadata.get_document_chunks(tenant["id"], document.id)
    assert all(value in "\n".join(record.text for record in records) for value in expected)
    query = "Invoice 00123 amount"
    hits = container.vectors.search(
        tenant["id"], container.embedder.embed([query])[0], top_k=10, query_text=query
    )
    assert all(value in "\n".join(hit.payload["content"] for hit in hits) for value in expected)


def test_scanned_pdf_preserves_transcribed_ids_amounts_and_page_order():

    document = FPDF()
    for color in ("white", "black"):
        document.add_page()
        document.image(Image.new("RGB", (100, 100), color), x=10, y=10, w=180, h=250)

    class Gateway:
        calls = 0

        def vision(self, image_bytes, prompt, mime):
            self.calls += 1
            return f"Invoice 00{self.calls} amount 007.00 due 2026-09-30"

    elements, summary = extract_document("scanned.pdf", bytes(document.output()), Gateway(), CFG)
    assert [element.text for element in elements] == [
        "Invoice 001 amount 007.00 due 2026-09-30",
        "Invoice 002 amount 007.00 due 2026-09-30",
    ]
    assert [element.meta["page"] for element in elements] == [1, 2]
    assert summary["extraction_status"] == "complete"


def test_wide_csv_preserves_every_header_cell_association_at_extraction():

    text = io.StringIO()
    writer = csv.writer(text)
    headers = [f"Column{index:02d}" for index in range(55)]
    values = [f"00{index:03d}" for index in range(55)]
    writer.writerow(headers)
    writer.writerow(values)
    elements, _ = extract_document("wide.csv", text.getvalue().encode(), NoGateway(), CFG)
    lines = elements[0].text.splitlines()
    assert [cell.strip() for cell in lines[0].strip("|").split("|")] == headers
    assert [cell.strip() for cell in lines[2].strip("|").split("|")] == values


def test_workbook_sparse_dimensions_are_bounded_before_iteration():

    workbook = Workbook()
    workbook.active["A1000"] = "00123"
    buffer = io.BytesIO()
    workbook.save(buffer)
    with pytest.raises(UnsafeContentError, match="row/cell limit"):
        extract_document(
            "sparse.xlsx",
            buffer.getvalue(),
            NoGateway(),
            CFG.model_copy(update={"max_workbook_cells": 50}),
        )


def test_csv_row_limit_never_indexes_a_truncated_prefix():

    with pytest.raises(UnsafeContentError, match="row limit"):
        extract_document(
            "rows.csv",
            b"ID\n001\n002\n003\n",
            NoGateway(),
            CFG.model_copy(update={"max_table_rows": 2}),
        )


def test_meaningful_heading_only_source_cannot_succeed_with_zero_chunks(container, tenant):

    document = Document(
        new_object_id(),
        tenant["id"],
        tenant["admin_id"],
        "text",
        "unused",
        "heading",
        "text/markdown",
        "heading.md",
        "private",
        [],
    )
    job = Job(new_object_id(), document.id, tenant["id"], "chunk", "running", 0)
    with pytest.raises(UnsafeContentError, match="no indexable chunks"):
        _stage_chunk(
            container,
            job,
            {
                "document": document,
                "elements": [Element("# Invoice 00123", "text", "markdown", "heading", 0)],
            },
        )
